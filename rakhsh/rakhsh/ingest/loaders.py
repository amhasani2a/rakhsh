"""Universal ingestion: any file / pasted text  ->  list[Chunk].

Supported: .txt .md (books, notes), pasted text, .csv/.tsv, .json/.jsonl,
.sqlite/.db, .xlsx, .pdf, .docx, .html. Books are chunked into passages that
remember their chapter; tables/databases become one chunk per row.
"""
from __future__ import annotations
import csv
import io
import json
import re
import sqlite3
import tempfile

from ..core.text import split_sentences
from .chunk import Chunk

_HEADING = re.compile(r"^\s*(#{1,6}\s+.+|(فصل|بخش|باب|دفتر|قسمت|گفتار|chapter|part|section)\b.{0,70}|[0-9۰-۹]{1,3}[\.\-\)]\s.{1,70})\s*$", re.I)


# ---------------------------------------------------------------- text / books
def decode(data: bytes) -> str:
    for enc in ("utf-8-sig", "utf-16", "cp1256", "latin-1"):
        try:
            t = data.decode(enc)
            if enc != "utf-16" or "\x00" not in t:
                return t
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", "ignore")


def chunk_text(text: str, source: str, max_words: int = 90, overlap: int = 1) -> list[Chunk]:
    chunks: list[Chunk] = []
    heading, buf, n_words, para_no = "", [], 0, 0

    def flush():
        nonlocal buf, n_words
        if buf:
            loc = f"{heading} · بند {len(chunks) + 1}" if heading else f"بند {len(chunks) + 1}"
            chunks.append(Chunk(source, loc, heading.lstrip("# ").strip(), " ".join(buf), sentences=list(buf)))
            buf = buf[-overlap:] if overlap and len(buf) > overlap else []
            n_words = sum(len(s.split()) for s in buf)

    for para in re.split(r"\n\s*\n|\r\n\s*\r\n", text.replace("\r", "")):
        lines = [l.strip() for l in para.split("\n") if l.strip()]
        if not lines:
            continue
        para_no += 1
        if len(lines) == 1 and len(lines[0]) < 90 and _HEADING.match(lines[0]):
            flush(); buf, n_words = [], 0; heading = lines[0].strip(); continue
        for s in split_sentences(" ".join(lines)):
            w = len(s.split())
            if n_words + w > max_words and n_words > 0:
                flush()
            buf.append(s); n_words += w
        if n_words > max_words * 0.6:
            flush()
    flush()
    return [c for c in chunks if len(c.text.split()) >= 2]


# ---------------------------------------------------------------- structured
def _row_chunk(source: str, locator: str, row: dict) -> Chunk:
    row = {str(k): ("" if v is None else str(v)) for k, v in row.items()}
    vals = [v for v in row.values() if v.strip()]
    title = next((v for v in vals if not re.fullmatch(r"[\d\.\-:/ ]+", v)), vals[0] if vals else "")[:80]
    text = " | ".join(f"{k}: {v}" for k, v in row.items() if v.strip())
    return Chunk(source, locator, title, text, kind="row", fields=row, sentences=[text])


def _flatten(obj, prefix="") -> dict:
    out = {}
    if isinstance(obj, dict):
        for k, v in obj.items():
            out.update(_flatten(v, f"{prefix}{k}."))
    elif isinstance(obj, list) and not all(isinstance(x, (dict, list)) for x in obj):
        out[prefix.rstrip(".")] = ", ".join(map(str, obj))
    elif isinstance(obj, list):
        for i, v in enumerate(obj):
            out.update(_flatten(v, f"{prefix}{i}."))
    else:
        out[prefix.rstrip(".")] = obj
    return out


def load_csv(data: bytes, source: str) -> list[Chunk]:
    text = decode(data)
    try:
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel_tab if source.endswith(".tsv") else csv.excel
    return [_row_chunk(source, f"ردیف {i + 1}", r) for i, r in enumerate(csv.DictReader(io.StringIO(text), dialect=dialect))]


def load_json(data: bytes, source: str) -> list[Chunk]:
    text = decode(data).strip()
    try:
        obj = json.loads(text)
    except json.JSONDecodeError:
        obj = [json.loads(l) for l in text.splitlines() if l.strip()]
    tables = {}
    if isinstance(obj, list):
        tables[""] = obj
    elif isinstance(obj, dict):
        lists = {k: v for k, v in obj.items() if isinstance(v, list) and v and isinstance(v[0], dict)}
        tables = lists or {"": [obj]}
    out = []
    for name, rows in tables.items():
        for i, r in enumerate(rows):
            r = _flatten(r) if isinstance(r, (dict, list)) else {"value": r}
            out.append(_row_chunk(source, f"{name + ' · ' if name else ''}رکورد {i + 1}", r))
    return out


def load_sqlite(data: bytes, source: str, max_rows: int = 50000) -> list[Chunk]:
    with tempfile.NamedTemporaryFile(suffix=".db") as f:
        f.write(data); f.flush()
        con = sqlite3.connect(f.name); con.row_factory = sqlite3.Row
        out = []
        for (t,) in con.execute("SELECT name FROM sqlite_master WHERE type='table' AND name NOT LIKE 'sqlite_%'"):
            for i, r in enumerate(con.execute(f'SELECT * FROM "{t}" LIMIT {max_rows}')):
                out.append(_row_chunk(source, f"جدول {t} · ردیف {i + 1}", dict(r)))
        con.close()
    return out


def load_xlsx(data: bytes, source: str) -> list[Chunk]:
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True); out = []
    for ws in wb.worksheets:
        rows = ws.iter_rows(values_only=True)
        header = [str(h) if h is not None else f"col{i}" for i, h in enumerate(next(rows, []))]
        for i, r in enumerate(rows):
            if any(v is not None for v in r):
                out.append(_row_chunk(source, f"شیت {ws.title} · ردیف {i + 1}", dict(zip(header, r))))
    return out


def load_pdf(data: bytes, source: str, max_words: int) -> list[Chunk]:
    import pypdf
    out = []
    for p, page in enumerate(pypdf.PdfReader(io.BytesIO(data)).pages):
        for c in chunk_text(page.extract_text() or "", source, max_words):
            c.locator = f"صفحه {p + 1} · {c.locator}"; out.append(c)
    return out


def load_docx(data: bytes, source: str, max_words: int) -> list[Chunk]:
    import docx
    d = docx.Document(io.BytesIO(data))
    parts = []
    for p in d.paragraphs:
        if p.text.strip():
            parts.append(("# " if p.style.name.lower().startswith(("heading", "title")) else "") + p.text)
    out = chunk_text("\n\n".join(parts), source, max_words)
    for ti, t in enumerate(d.tables):
        rows = [[c.text for c in r.cells] for r in t.rows]
        if len(rows) > 1:
            out += [_row_chunk(source, f"جدول {ti + 1} · ردیف {i + 1}", dict(zip(rows[0], r))) for i, r in enumerate(rows[1:])]
    return out


def load_html(data: bytes, source: str, max_words: int) -> list[Chunk]:
    t = re.sub(r"(?is)<(script|style).*?</\1>", " ", decode(data))
    t = re.sub(r"(?i)<(h[1-6])[^>]*>(.*?)</\1>", r"\n\n# \2\n\n", t)
    t = re.sub(r"(?i)<(br|/p|/div|/li)[^>]*>", "\n\n", t)
    return chunk_text(re.sub(r"<[^>]+>", " ", t), source, max_words)


